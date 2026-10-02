"""선택 월 손익표 엑셀 내보내기 (FR-23). 웹 버전과 같은 시트 구성."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from .categories import flatten
from .fmt import EMPTY

#: (지표 키, 이름, 단위) — 화면 열 순서
COLUMNS: tuple[tuple[str, str, str], ...] = (
    ("salesVolume", "판매량", "천톤"),
    ("revenue", "매출", "억원"),
    ("operatingProfit", "영업이익", "억원"),
    ("operatingMargin", "이익률", "%"),
    ("ordinaryProfit", "경상이익", "억원"),
    ("ordinaryMargin", "경상이익률", "%"),
)

NUMBER_FORMAT = {
    "salesVolume": "#,##0.0",
    "revenue": "#,##0",
    "operatingProfit": "#,##0",
    "ordinaryProfit": "#,##0",
    "operatingMargin": "0.0",
    "ordinaryMargin": "0.0",
}


def month_label(month: str) -> str:
    year, m = month.split("-")
    return f"{int(year)}년 {int(m)}월"


def build_rows(month: str, metrics: dict[str, dict]) -> list[list]:
    """[제목], [], [헤더], 데이터 행… 숫자 셀은 (값, 서식) 튜플."""
    rows: list[list] = [[f"{month_label(month)} 손익 실적"], [], ["구분", *(f"{n}({u})" for _, n, u in COLUMNS)]]
    for cat in flatten():
        m = metrics.get(cat.id, {})
        row: list = ["    " * cat.level + cat.name]
        for key, _, _ in COLUMNS:
            value = m.get(key)
            row.append(EMPTY if value is None else (value, NUMBER_FORMAT[key]))
        rows.append(row)
    return rows


def to_xlsx(month: str, metrics: dict[str, dict]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = month_label(month)
    for r, row in enumerate(build_rows(month, metrics), start=1):
        for c, cell in enumerate(row, start=1):
            target = sheet.cell(row=r, column=c)
            if isinstance(cell, tuple):
                target.value, target.number_format = cell
            else:
                target.value = cell
    sheet["A1"].font = Font(bold=True, size=13)
    for c in range(1, 8):
        sheet.cell(row=3, column=c).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 18
    for letter in "BCDEFG":
        sheet.column_dimensions[letter].width = 14
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()
